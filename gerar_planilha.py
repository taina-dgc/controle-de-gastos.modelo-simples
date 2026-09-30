# -*- coding: utf-8 -*-
"""
Gera a planilha "Contas Mensais" (Excel .xlsx) a partir de dois arquivos:

  config.json  -> as REGRAS (categorias, formas de pagamento, cartão, nº de meses)
  dados        -> os VALORES (salário, contas, parcelamentos, gastos do mês)

Uso:
  python gerar_planilha.py                                   # usa dados/exemplo.json
  python gerar_planilha.py --dados ../privado/meus_dados.json --saida Minhas_Contas.xlsx

Abas geradas:
  1. Contas Mensais  – parcelas, dívidas e contas fixas por mês + quanto sobra do salário
  2. Gastos do Mês   – gastos do dia a dia com limite por semana (ligado à fatura)
  3. Como usar       – instruções
"""
import argparse
import datetime as dt
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

# --------------------------------------------------------------------------- estilo
NAVY, PINK, YEL = "1F3864", "EAD1DC", "FFF2CC"
THIN = Side(style="thin", color="D9D9D9")
BORDA = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MOEDA = '"R$" #,##0.00;[Red]-"R$" #,##0.00;"–"'
MES_FMT = "mmm/yy"


def fonte(**k):
    return Font(name="Arial", size=k.pop("size", 10), **k)


def preencher(cor):
    return PatternFill("solid", fgColor=cor)


def mes(txt):
    """'2026-10' -> date(2026, 10, 1)"""
    if not txt:
        return None
    a, m = map(int, str(txt)[:7].split("-"))
    return dt.date(a, m, 1)


def proxima_fatura(hoje, dia_venc):
    """Mês da próxima fatura a pagar (se já passou o vencimento, é a do mês seguinte)."""
    a, m = hoje.year, hoje.month + (1 if hoje.day >= dia_venc else 0)
    if m > 12:
        a, m = a + 1, 1
    return dt.date(a, m, 1)


def idx(ref):
    """Fórmula que transforma uma data em 'nº do mês' (ano*12+mês)."""
    return f"(YEAR({ref})*12+MONTH({ref}))"


# --------------------------------------------------------------------------- planilha
class Gerador:
    def __init__(self, config, dados):
        self.cfg, self.dados = config, dados
        self.wb = Workbook()
        self.nm = config.get("meses_na_tabela", 15)
        self.linhas_contas = max(config.get("linhas_de_contas", 40), len(dados.get("contas", [])) + 10)
        self.linhas_gastos = config.get("linhas_de_gastos", 200)

    def nome(self, n, ref):
        self.wb.defined_names[n] = DefinedName(n, attr_text=ref)

    def lista(self, ws, formula, faixa, aviso=None, estrito=True):
        d = DataValidation(type="list", formula1=formula, allow_blank=True)
        d.showErrorMessage = estrito
        d.error = "Escolha um item da lista."
        if aviso:
            d.showInputMessage, d.prompt = True, aviso
        ws.add_data_validation(d)
        d.add(faixa)

    # ------------------------------------------------------------- aba 1: Contas Mensais
    def contas_mensais(self):
        ws = self.wb.active
        ws.title = "Contas Mensais"
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = "C00000"
        S = "'Contas Mensais'!"
        NM, HR, G0 = self.nm, 10, 11
        GL = G0 + self.linhas_contas - 1
        MC = [CL(11 + i) for i in range(NM)]
        ULT = MC[-1]
        self.cm = dict(S=S, HR=HR, G0=G0, GL=GL, MC=MC)

        # título
        ws.merge_cells(f"A1:{ULT}1")
        ws["A1"] = "CONTAS MENSAIS — PARCELAS, DÍVIDAS E QUANTO SOBRA DO SALÁRIO"
        ws["A1"].font = fonte(size=15, bold=True, color="FFFFFF")
        ws["A1"].alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[1].height = 30
        for c in range(1, 11 + NM):
            ws.cell(1, c).fill = preencher(NAVY)
        ws.merge_cells(f"A2:{ULT}2")
        ws["A2"] = ("Cadastre cada conta UMA vez: o que é, a quem, forma de pagamento, valor TOTAL, nº de parcelas e o mês "
                    "da 1ª parcela (no cartão = mês da fatura). Conta fixa: deixe Parcelas vazio e use o valor mensal. "
                    "Escolha o mês em B4 para ver a fatura e a sobra daquele mês.")
        ws["A2"].font = fonte(size=9, italic=True, color="595959")
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="center", indent=1)
        ws.row_dimensions[2].height = 30

        # entradas do topo
        ws["A4"], ws["D4"] = "MÊS:", "ÚLTIMO SALÁRIO:"
        for c in ("A4", "D4"):
            ws[c].font = fonte(bold=True, size=11)
            ws[c].alignment = Alignment(horizontal="right")
        ws["B4"] = f"=$K${HR}"
        ws["B4"].number_format = "mmm/yyyy"
        ws["E4"] = self.dados.get("salario", 0)
        ws["E4"].number_format = MOEDA
        ws["F4"] = "◄ atualize todo mês com o salário que caiu"
        ws["F4"].font = fonte(size=8, italic=True, color="595959")
        for c in ("B4", "E4"):
            ws[c].fill = preencher(YEL)
            ws[c].font = fonte(bold=True, size=12, color="0000FF")
            ws[c].border = BORDA
            ws[c].alignment = Alignment(horizontal="center")
        self.nome("sel", S + "$B$4")
        self.nome("sal", S + "$E$4")
        self.lista(ws, f"{S}$K${HR}:${ULT}${HR}", "B4", "Escolha o mês.")

        # cabeçalho da tabela
        cols = [("STATUS", 14), ("GASTO", 30), ("A QUEM DEVO", 18), ("FORMA", 9), ("VALOR TOTAL (fixo: valor mensal)", 13),
                ("PARCELAS (vazio = todo mês)", 9), ("1º MÊS", 9), ("PARCELA", 11), ("FALTA PAGAR", 12), ("ÚLTIMA PARCELA", 10)]
        for i, (t, w) in enumerate(cols):
            c = ws.cell(HR, i + 1, t)
            ws.column_dimensions[CL(i + 1)].width = w
            c.font = fonte(bold=True, color="FFFFFF", size=9)
            c.fill = preencher("7F7F7F" if i in (0, 7, 8, 9) else NAVY)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = BORDA
        ws.row_dimensions[HR].height = 40
        inicio = mes(self.dados.get("primeiro_mes_da_tabela")) or proxima_fatura(dt.date.today(), self.cfg["cartao"]["dia_vencimento"])
        for i, col in enumerate(MC):
            ws.column_dimensions[col].width = 12
            ws[f"{col}{HR}"] = inicio if i == 0 else f"=DATE(YEAR({MC[i-1]}{HR}),MONTH({MC[i-1]}{HR})+1,1)"
            c = ws[f"{col}{HR}"]
            c.number_format = MES_FMT
            c.font = fonte(bold=True, color="FFFFFF", size=9)
            c.fill = preencher("C00000")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = BORDA
        ws["K9"] = "◄ 1º mês da tabela (troque a data em K10 para avançar)"
        ws["K9"].font = fonte(size=8, italic=True, color="595959")

        # linhas de contas
        for r in range(G0, GL + 1):
            for c in range(1, 11 + NM):
                x = ws.cell(r, c)
                x.border = BORDA
                x.fill = preencher("FFFFFF" if 2 <= c <= 7 else "F2F2F2")
                x.font = fonte(size=9, color="0000FF" if 2 <= c <= 7 else "000000")
                if 2 <= c <= 7:
                    x.protection = Protection(locked=False)
            for c in "EHI":
                ws[f"{c}{r}"].number_format = MOEDA
            for c in "GJ":
                ws[f"{c}{r}"].number_format = MES_FMT
            for c in "AFGJ":
                ws[f"{c}{r}"].alignment = Alignment(horizontal="center")
            ws[f"H{r}"] = f'=IF(OR(B{r}="",E{r}=""),"",IF(N(F{r})=0,E{r},E{r}/F{r}))'
            ws[f"J{r}"] = f'=IF(OR(H{r}="",G{r}="",N(F{r})=0),"",DATE(YEAR(G{r}),MONTH(G{r})+F{r}-1,1))'
            ws[f"I{r}"] = f'=IF(J{r}="",0,H{r}*MAX(0,{idx(f"J{r}")}-MAX({idx(f"G{r}")},{idx(f"$K${HR}")})+1))'
            ws[f"A{r}"] = (f'=IF(B{r}="","",IF(G{r}="","Falta 1º mês",IF(N(F{r})=0,"Fixo",IF(J{r}<$K${HR},"Quitado",'
                           f'IF(J{r}=sel,"Última parcela","Ativo")))))')
            for col in MC:
                ws[f"{col}{r}"] = (f'=IF(OR($H{r}="",$G{r}=""),"",IF(AND({col}${HR}>=$G{r},OR(N($F{r})=0,{col}${HR}<=$J{r})),$H{r},""))')
                ws[f"{col}{r}"].number_format = MOEDA
            ws[f"Z{r}"] = f'=IF(OR(J{r}="",D{r}="Caixinha"),0,H{r}*MAX(0,{idx(f"J{r}")}-MAX({idx(f"G{r}")},{idx("sel")})+1))'
        ws.column_dimensions["Z"].hidden = True
        formas = ",".join(self.cfg["formas_contas"])
        self.lista(ws, f'"{formas}"', f"D{G0}:D{GL}", "Cartão = entra na fatura. Pix/Boleto = sai da conta. Caixinha = meta de economia.")
        self.lista(ws, f"{S}$K${HR}:${ULT}${HR}", f"G{G0}:G{GL}",
                   "Mês da 1ª parcela. No cartão, é o mês da FATURA. Pode digitar um mês anterior (ex.: 01/05/2026).", estrito=False)

        # cores
        for t, f, fc in [('"Quitado"', "B6D7A8", "1E5631"), ('"Ativo"', YEL, "7F6000"), ('"Fixo"', "DDEBF7", "1F3864"),
                         ('"Última parcela"', "F8CBAD", "843C0C"), ('"Falta 1º mês"', "F8D7DA", "9C0006")]:
            ws.conditional_formatting.add(f"A{G0}:A{GL}", CellIsRule(operator="equal", formula=[t], fill=preencher(f), font=Font(color=fc, bold=True)))
        grade = f"K{G0}:{ULT}{GL}"
        ws.conditional_formatting.add(grade, FormulaRule(formula=[f'AND(K{G0}<>"",K${HR}=$J{G0})'], fill=preencher(PINK), font=Font(bold=True)))
        ws.conditional_formatting.add(grade, FormulaRule(formula=[f"K${HR}=sel"], fill=preencher("FFF9DB")))
        ws.conditional_formatting.add(f"K{HR}:{ULT}{HR}", FormulaRule(formula=[f"K${HR}=sel"], fill=preencher("FFC000"), font=Font(color="000000", bold=True)))
        ws.conditional_formatting.add(f"A{G0}:{ULT}{GL}", FormulaRule(formula=[f'$A{G0}="Quitado"'], font=Font(color="A6A6A6")))

        # totais
        T = GL + 2
        self.cm["T"] = T
        linhas = [("FATURA DO CARTÃO", lambda c: f'=SUMIFS({c}{G0}:{c}{GL},$D${G0}:$D${GL},"Cartão")', "DDEBF7"),
                  ("PIX / BOLETO (pessoas, aluguel…)", lambda c: f'=SUMIFS({c}{G0}:{c}{GL},$D${G0}:$D${GL},"<>Cartão")-SUMIFS({c}{G0}:{c}{GL},$D${G0}:$D${GL},"Caixinha")', "FCE4D6"),
                  ("GUARDAR NA CAIXINHA", lambda c: f'=SUMIFS({c}{G0}:{c}{GL},$D${G0}:$D${GL},"Caixinha")', "E2EFDA"),
                  ("TOTAL MENSAL", lambda c: f"=SUM({c}{G0}:{c}{GL})", "F4B6B6"),
                  ("ÚLTIMO SALÁRIO", lambda c: "=sal", "E2EFDA"),
                  ("SOBRA PARA O DIA A DIA", lambda c: f"={c}{T+4}-{c}{T+3}", "FFF2CC"),
                  ("SOBRA POR DIA", lambda c: f"={c}{T+5}/DAY(DATE(YEAR({c}${HR}),MONTH({c}${HR})+1,0))", "FFF2CC"),
                  ("% DO SALÁRIO COMPROMETIDO", lambda c: f"=IF(sal=0,0,{c}{T+3}/sal)", "FFFFFF")]
        for i, (rot, fn, cor) in enumerate(linhas):
            r = T + i
            ws.merge_cells(f"A{r}:J{r}")
            ws[f"A{r}"] = rot
            for c in range(1, 11 + NM):
                x = ws.cell(r, c)
                x.fill, x.border = preencher(cor), BORDA
                x.font = fonte(bold=True, size=10 if i in (3, 5) else 9)
            ws[f"A{r}"].alignment = Alignment(horizontal="right", indent=1)
            for col in MC:
                ws[f"{col}{r}"] = fn(col)
                ws[f"{col}{r}"].number_format = "0%" if i == 7 else MOEDA
        for rr in (T + 5, T + 6):
            ws.conditional_formatting.add(f"K{rr}:{ULT}{rr}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True), fill=preencher("F8D7DA")))
        self.nome("tot", f"{S}$K${T}:${ULT}${T+7}")
        self.nome("hdr", f"{S}$K${HR}:${ULT}${HR}")

        # cartões do topo
        V = lambda k: f"=INDEX(tot,{k},MATCH(sel,hdr,0))"
        cards = [("FATURA DO CARTÃO", V(1), "2F5597"), ("PIX / BOLETO", V(2), "C65911"), ("CAIXINHA", V(3), "548235"),
                 ("TOTAL DO MÊS", V(4), "C00000"), ("SOBRA DO SALÁRIO", V(6), NAVY), ("POR DIA", V(7), NAVY),
                 ("AINDA DEVO (do mês em diante)", f"=SUM(Z{G0}:Z{GL})", "C00000")]
        for n, (rot, f, cor) in enumerate(cards):
            a, b = CL(1 + 2 * n), CL(2 + 2 * n)
            for r in (6, 7, 8):
                ws.merge_cells(f"{a}{r}:{b}{r}")
                for cc in (1 + 2 * n, 2 + 2 * n):
                    x = ws.cell(r, cc)
                    x.fill = preencher("F7F9FC")
                    x.border = Border(left=THIN, right=THIN, top=Side(style="thick", color=cor) if r == 6 else None, bottom=THIN if r == 8 else None)
            ws[f"{a}6"], ws[f"{a}7"] = rot, f
            ws[f"{a}6"].font = fonte(size=8, bold=True, color="595959")
            ws[f"{a}7"].font = fonte(size=14, bold=True, color=cor)
            ws[f"{a}7"].number_format = MOEDA
            for r in (6, 7, 8):
                ws[f"{a}{r}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws["A8"] = '=TEXT(INDEX(tot,8,MATCH(sel,hdr,0)),"0%")&" do salário"'
        ws["C8"] = f'="vence dia {self.cfg["cartao"]["dia_vencimento"]:02d}"'
        ws["I8"] = '=IF(INDEX(tot,6,MATCH(sel,hdr,0))<0,"o mês não fecha","para mercado, transporte, lazer")'
        ws["K8"] = '=DAY(DATE(YEAR(sel),MONTH(sel)+1,0))&" dias no mês"'
        ws["M8"] = f'=COUNTIF(J{G0}:J{GL},sel)&" parcela(s) terminam neste mês"'
        for c in ("A8", "C8", "I8", "K8", "M8"):
            ws[c].font = fonte(size=8, color="595959")
        ws.row_dimensions[7].height = 30
        ws.conditional_formatting.add("I7:L7", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))

        # gráfico
        g = BarChart()
        g.type, g.grouping, g.overlap = "col", "stacked", 100
        g.title, g.height, g.width = "Quanto sai por mês x salário", 8, 26
        g.add_data(Reference(ws, min_col=11, max_col=10 + NM, min_row=T, max_row=T + 2), from_rows=True, titles_from_data=False)
        for s, (n, cor) in zip(g.series, [("Fatura do cartão", "5B7DB1"), ("Pix / boleto", "E0A33B"), ("Caixinha", "70AD47")]):
            s.tx = SeriesLabel(v=n)
            s.graphicalProperties.solidFill = cor
        g.set_categories(Reference(ws, min_col=11, max_col=10 + NM, min_row=HR))
        ln = LineChart()
        ln.add_data(Reference(ws, min_col=11, max_col=10 + NM, min_row=T + 4), from_rows=True, titles_from_data=False)
        ln.series[0].tx = SeriesLabel(v="Salário")
        ln.series[0].graphicalProperties.line.solidFill = "C00000"
        ln.series[0].smooth = False
        g += ln
        g.y_axis.numFmt, g.legend.position = "#,##0", "b"
        g.y_axis.delete = g.x_axis.delete = False
        ws.add_chart(g, f"A{T+10}")
        ws.freeze_panes = f"C{G0}"

        # contas
        for i, c in enumerate(self.dados.get("contas", [])):
            r = G0 + i
            ws[f"B{r}"], ws[f"C{r}"] = c.get("gasto"), c.get("a_quem")
            ws[f"D{r}"], ws[f"E{r}"] = c.get("forma"), c.get("valor_total")
            if c.get("parcelas"):
                ws[f"F{r}"] = c["parcelas"]
            ws[f"G{r}"] = mes(c.get("primeiro_mes"))
        # linha automática: gastos do mês no cartão (vem da aba Gastos do Mês)
        self.linha_link = G0 + len(self.dados.get("contas", []))

    # ------------------------------------------------------------- aba 2: Gastos do Mês
    def gastos_do_mes(self):
        g = self.wb.create_sheet("Gastos do Mês", 1)
        g.sheet_view.showGridLines = False
        g.sheet_properties.tabColor = "70AD47"
        for k, v in {"A": 2, "B": 12, "C": 30, "D": 20, "E": 15, "F": 13, "G": 10, "H": 3, "I": 12, "J": 11, "K": 11, "L": 13, "M": 13, "N": 13, "O": 13, "P": 16}.items():
            g.column_dimensions[k].width = v
        g.merge_cells("B1:P1")
        g["B1"] = "GASTOS DO MÊS — LIMITE SEMANAL"
        g["B1"].font = fonte(size=15, bold=True, color="FFFFFF")
        for c in range(1, 17):
            g.cell(1, c).fill = preencher(NAVY)
        g.row_dimensions[1].height = 30
        g.merge_cells("B2:P2")
        g["B2"] = ("Anote aqui TODO gasto do dia a dia. Os limites das semanas são editáveis (células amarelas). "
                   "Os gastos no CARTÃO entram sozinhos na fatura do mês seguinte, na aba Contas Mensais.")
        g["B2"].font = fonte(size=9, italic=True, color="595959")
        g["B2"].alignment = Alignment(wrap_text=True, vertical="center")

        def entrada(cel, val, fmt):
            c = g[cel]
            c.value, c.number_format = val, fmt
            c.fill, c.border = preencher(YEL), BORDA
            c.font = fonte(bold=True, size=12, color="0000FF")
            c.alignment = Alignment(horizontal="center")
            c.protection = Protection(locked=False)

        gm = self.dados.get("gastos_do_mes", {})
        g["B4"], g["D4"] = "Mês:", "Limite do mês:"
        for c in ("B4", "D4"):
            g[c].font = fonte(bold=True, size=11)
            g[c].alignment = Alignment(horizontal="right")
        entrada("C4", mes(gm.get("mes")) or dt.date.today().replace(day=1), "mmmm/yyyy")
        entrada("E4", gm.get("limite_do_mes", 0), MOEDA)
        g["F4"] = "◄ troque o mês e o limite quando começar um mês novo"
        g["F4"].font = fonte(size=8, italic=True, color="595959")
        g["C5"] = "=DATE(YEAR(gm_Mes),MONTH(gm_Mes)+1,0)"
        g["C5"].number_format = "dd/mm/yyyy"
        g["C5"].font = fonte(size=8, color="A6A6A6")
        self.nome("gm_Mes", "'Gastos do Mês'!$C$4")
        self.nome("gm_Lim", "'Gastos do Mês'!$E$4")
        self.nome("gm_Fim", "'Gastos do Mês'!$C$5")

        # lista de gastos
        L0, LL = 26, 26 + self.linhas_gastos - 1
        g["B24"] = "ANOTE SEUS GASTOS"
        g["B24"].font = fonte(bold=True, size=11, color=NAVY)
        for i, t in enumerate(["Data", "Descrição", "Categoria", "Pago com", "Valor", "Semana"]):
            c = g.cell(25, 2 + i, t)
            c.font = fonte(bold=True, color="FFFFFF", size=9)
            c.fill = preencher("7F7F7F" if i == 5 else "548235")
            c.alignment, c.border = Alignment(horizontal="center"), BORDA
        for r in range(L0, LL + 1):
            for c in range(2, 8):
                x = g.cell(r, c)
                x.border = BORDA
                x.font = fonte(size=9, color="000000" if c == 7 else "0000FF")
                x.fill = preencher("F2F2F2" if c == 7 else "FFFFFF")
                if c < 7:
                    x.protection = Protection(locked=False)
            g[f"B{r}"].number_format = "dd/mm"
            g[f"F{r}"].number_format = MOEDA
            g[f"G{r}"] = f'=IF(B{r}="","",IF(OR(B{r}<gm_Mes,B{r}>gm_Fim),"fora do mês",ROUNDUP(DAY(B{r})/7,0)))'
            g[f"G{r}"].alignment = Alignment(horizontal="center")
        cats = self.cfg["categorias_gastos"]
        self.lista(g, '"' + ",".join(cats) + '"', f"D{L0}:D{LL}", estrito=False)
        self.lista(g, '"' + ",".join(self.cfg["formas_gastos"]) + '"', f"E{L0}:E{LL}")
        for n, col in zip(["gm_Data", "gm_Cat", "gm_Pag", "gm_Val", "gm_Sem"], "BDEFG"):
            self.nome(n, f"'Gastos do Mês'!${col}${L0}:${col}${LL}")
        g.auto_filter.ref = f"B25:G{LL}"
        for i, x in enumerate(gm.get("gastos", [])):
            r = L0 + i
            g[f"B{r}"] = dt.date.fromisoformat(x["data"])
            g[f"C{r}"], g[f"D{r}"], g[f"E{r}"], g[f"F{r}"] = x.get("descricao"), x.get("categoria"), x.get("pago_com"), x.get("valor")

        # cartões
        cm = self.cm
        TOT = 'SUMIFS(gm_Val,gm_Data,">="&gm_Mes,gm_Data,"<="&gm_Fim)'
        HOJE = "IF(TODAY()<gm_Mes,gm_Mes,TODAY())"
        DIAS = f"MAX(1,gm_Fim-{HOJE}+1)"
        U = cm["MC"][-1]
        SOBRA = (f"IFERROR(INDEX({cm['S']}$K${cm['T']+5}:${U}${cm['T']+5},MATCH(DATE(YEAR(gm_Mes),MONTH(gm_Mes)+1,1),"
                 f"{cm['S']}$K${cm['HR']}:${U}${cm['HR']},0)),0)")
        NCART = 'SUMIFS(gm_Val,gm_Pag,"<>Cartão",gm_Data,">="&gm_Mes,gm_Data,"<="&gm_Fim)'
        cards = [("LIMITE DO MÊS", "=gm_Lim", NAVY, None),
                 ("JÁ GASTEI", f"={TOT}", "C00000", f'=IF(gm_Lim=0,"",TEXT({TOT}/gm_Lim,"0%")&" do limite")'),
                 ("AINDA POSSO GASTAR", f"=gm_Lim-{TOT}", "548235", f'=IF(AND(TODAY()>=gm_Mes,TODAY()<=gm_Fim),"faltam "&{DIAS}&" dias","")'),
                 ("POR DIA (daqui até o fim do mês)", f"=IF(TODAY()>gm_Fim,0,(gm_Lim-{TOT})/{DIAS})", "548235", None),
                 ("RESPIRO PREVISTO NO SALÁRIO SEGUINTE", f"={SOBRA}-{NCART}", "2F5597", '="sobra da aba Contas Mensais − gastos no débito/Pix"'),
                 ("RESPIRO SE USAR TODO O LIMITE", f"={SOBRA}-{NCART}-MAX(0,gm_Lim-{TOT})", "2F5597", None)]
        for (rot, f, cor, sub), (c1, c2) in zip(cards, [(2, 3), (4, 5), (6, 7), (9, 10), (11, 13), (14, 16)]):
            a, b = CL(c1), CL(c2)
            for r in (7, 8, 9):
                g.merge_cells(f"{a}{r}:{b}{r}")
                for cc in range(c1, c2 + 1):
                    x = g.cell(r, cc)
                    x.fill = preencher("F7F9FC")
                    x.border = Border(left=THIN, right=THIN, top=Side(style="thick", color=cor) if r == 7 else None, bottom=THIN if r == 9 else None)
            g[f"{a}7"], g[f"{a}8"] = rot, f
            if sub:
                g[f"{a}9"] = sub
            g[f"{a}7"].font = fonte(size=8, bold=True, color="595959")
            g[f"{a}8"].font = fonte(size=15, bold=True, color=cor)
            g[f"{a}8"].number_format = MOEDA
            g[f"{a}9"].font = fonte(size=8, color="595959")
            for r in (7, 8, 9):
                g[f"{a}{r}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        g.row_dimensions[7].height, g.row_dimensions[8].height = 26, 30
        for c in ("F8", "I8", "K8", "N8"):
            g.conditional_formatting.add(c, CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True)))

        # semanas
        g["B11"] = "LIMITE POR SEMANA"
        g["B11"].font = fonte(bold=True, size=11, color=NAVY)
        pos = [2, 3, 4, 5, 6, 7, 9, 10, 11]
        g.merge_cells("K12:L12")
        for t, c in zip(["Semana", "De", "Até", "Limite (editável)", "Gastei", "Sobrou na semana", "Saldo acumulado do mês", "% usado", "Situação"], pos):
            x = g.cell(12, c, t)
            x.font = fonte(bold=True, color="FFFFFF", size=9)
            x.fill = preencher("548235" if c == 5 else "7F7F7F")
            x.alignment, x.border = Alignment(horizontal="center", vertical="center", wrap_text=True), BORDA
        g.row_dimensions[12].height = 30
        lim_sem = gm.get("limites_por_semana", [])
        for k in range(1, 6):
            r = 12 + k
            g[f"B{r}"] = f"Semana {k}"
            ini = 7 * (k - 1) + 1
            g[f"C{r}"] = f'=IF(DATE(YEAR(gm_Mes),MONTH(gm_Mes),{ini})>gm_Fim,"",DATE(YEAR(gm_Mes),MONTH(gm_Mes),{ini}))'
            g[f"D{r}"] = f'=IF(C{r}="","",MIN(gm_Fim,C{r}+6))'
            g[f"E{r}"] = lim_sem[k - 1] if k < 5 and k - 1 < len(lim_sem) else ('=IF(C17="",0,MAX(0,gm_Lim-SUM(E13:E16)))' if k == 5 else 0)
            g[f"F{r}"] = f"=SUMIFS(gm_Val,gm_Sem,{k})"
            g[f"G{r}"] = f'=IF(C{r}="","",E{r}-F{r})'
            g[f"I{r}"] = f'=IF(C{r}="","",SUM($E$13:E{r})-SUM($F$13:F{r}))'
            g[f"J{r}"] = f'=IF(OR(C{r}="",N(E{r})=0),"",F{r}/E{r})'
            g.merge_cells(f"K{r}:L{r}")
            g[f"K{r}"] = (f'=IF(C{r}="","",IF(TODAY()<C{r},"ainda não começou",IF(F{r}>E{r},"passou do limite",'
                          f'IF(TODAY()<=D{r},"em andamento","dentro do limite ✔"))))')
            for c in pos + [12]:
                x = g.cell(r, c)
                x.border, x.font, x.alignment = BORDA, fonte(size=9), Alignment(horizontal="center")
            for c in "EFGI":
                g[f"{c}{r}"].number_format = MOEDA
            g[f"C{r}"].number_format = g[f"D{r}"].number_format = "dd/mm"
            g[f"J{r}"].number_format = "0%"
            e = g[f"E{r}"]
            e.fill, e.font, e.protection = preencher(YEL), fonte(size=10, bold=True, color="0000FF"), Protection(locked=False)
        g["B18"] = "Mês"
        g["E18"], g["F18"], g["G18"], g["J18"] = "=SUM(E13:E17)", "=SUM(F13:F17)", "=E18-F18", '=IF(E18=0,"",F18/E18)'
        g.merge_cells("K18:L18")
        g["K18"] = '=IF(ROUND(E18,2)<>ROUND(gm_Lim,2),"⚠ semanas somam "&TEXT(E18,"0")&" (limite "&TEXT(gm_Lim,"0")&")","semanas = limite do mês ✔")'
        for c in pos + [12]:
            x = g.cell(18, c)
            x.border, x.font, x.fill, x.alignment = BORDA, fonte(size=9, bold=True), preencher("E2EFDA"), Alignment(horizontal="center")
        for c in "EFG":
            g[f"{c}18"].number_format = MOEDA
        g["J18"].number_format = "0%"
        g.merge_cells("B19:P20")
        g["B19"] = ("A semana 5 recebe automaticamente o que faltar para completar o limite do mês (pode digitar outro valor). "
                    "Se passar do limite numa semana, o 'Saldo acumulado' mostra quanto compensar nas próximas.")
        g["B19"].font = fonte(size=8, italic=True, color="595959")
        g["B19"].alignment = Alignment(wrap_text=True, vertical="top")
        g.conditional_formatting.add("J13:J18", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="70AD47"))
        g.conditional_formatting.add("G13:I18", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000", bold=True), fill=preencher("F8D7DA")))
        for t, f, fc in [('"passou do limite"', "F8D7DA", "9C0006"), ('"dentro do limite ✔"', "D4EDDA", "1E5631"), ('"em andamento"', "FFF3CD", "7F6000")]:
            g.conditional_formatting.add("K13:K17", CellIsRule(operator="equal", formula=[t], font=Font(color=fc, bold=True), fill=preencher(f)))

        # por categoria e forma de pagamento
        g["I24"] = "ONDE ESTOU GASTANDO"
        g["I24"].font = fonte(bold=True, size=11, color=NAVY)
        for t, c in zip(["Categoria", "", "Gasto", "% do total"], [9, 10, 11, 12]):
            x = g.cell(25, c, t)
            x.font, x.fill, x.alignment, x.border = fonte(bold=True, color="FFFFFF", size=9), preencher("7F7F7F"), Alignment(horizontal="center"), BORDA
        g.merge_cells("I25:J25")
        for i, cat in enumerate(cats + ["Total"]):
            r = 26 + i
            g.merge_cells(f"I{r}:J{r}")
            g[f"I{r}"] = cat
            g[f"K{r}"] = f"={TOT}" if cat == "Total" else f'=SUMIFS(gm_Val,gm_Cat,I{r},gm_Data,">="&gm_Mes,gm_Data,"<="&gm_Fim)'
            g[f"L{r}"] = f"=IF({TOT}=0,0,K{r}/{TOT})"
            for c in range(9, 13):
                x = g.cell(r, c)
                x.border, x.font = BORDA, fonte(size=9, bold=(cat == "Total"))
            g[f"K{r}"].number_format, g[f"L{r}"].number_format = MOEDA, "0%"
        g.conditional_formatting.add(f"L26:L{25+len(cats)}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="C0504D"))
        base = 28 + len(cats)
        g[f"I{base}"] = "Por forma de pagamento"
        g[f"I{base}"].font = fonte(bold=True, size=9, color=NAVY)
        for i, pg in enumerate(self.cfg["formas_gastos"]):
            r = base + 1 + i
            g.merge_cells(f"I{r}:J{r}")
            g[f"I{r}"] = pg
            g[f"K{r}"] = f'=SUMIFS(gm_Val,gm_Pag,I{r},gm_Data,">="&gm_Mes,gm_Data,"<="&gm_Fim)'
            g[f"K{r}"].number_format = MOEDA
            for c in range(9, 12):
                g.cell(r, c).border = BORDA
        g.freeze_panes = "A6"

        # liga os gastos no cartão à tabela de Contas Mensais
        ws = self.wb["Contas Mensais"]
        r = self.linha_link
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"F{r}"] = "Gastos do mês no cartão (aba Gastos do Mês)", "Cartão de crédito", "Cartão", 1
        ws[f"E{r}"] = (f"=SUMIFS('Gastos do Mês'!$F${L0}:$F${LL},'Gastos do Mês'!$E${L0}:$E${LL},\"Cartão\","
                       f"'Gastos do Mês'!$B${L0}:$B${LL},\">=\"&gm_Mes,'Gastos do Mês'!$B${L0}:$B${LL},\"<=\"&gm_Fim)")
        ws[f"G{r}"] = "=DATE(YEAR(gm_Mes),MONTH(gm_Mes)+1,1)"

    # ------------------------------------------------------------- aba 3: Como usar
    def como_usar(self):
        h = self.wb.create_sheet("Como usar")
        h.sheet_view.showGridLines = False
        h.column_dimensions["A"].width = 115
        txt = ["COMO USAR",
               "1. Todo mês, quando o salário cair, atualize o ÚLTIMO SALÁRIO (Contas Mensais, E4).",
               "2. Compra parcelada nova: numa linha branca de Contas Mensais, escreva o que é, a quem deve, forma, VALOR TOTAL, nº de PARCELAS e o 1º MÊS.",
               f"   No cartão, o 1º mês é o da FATURA (a fatura fecha {self.cfg['cartao']['dias_entre_fechamento_e_vencimento']} dias antes do vencimento, dia {self.cfg['cartao']['dia_vencimento']}).",
               "3. Conta fixa (aluguel, assinatura, academia): deixe PARCELAS vazio e coloque o valor mensal.",
               "4. Meta de economia: forma 'Caixinha'. Entra no total do mês, mas não conta como dívida.",
               "5. Escolha o MÊS em B4 para ver a fatura, o total, a sobra do salário e quanto dá por dia.",
               "6. Gastos do dia a dia: anote na aba 'Gastos do Mês'. Os feitos no cartão entram sozinhos na fatura do mês seguinte.",
               "7. Mês novo: na aba Gastos do Mês, troque o mês (C4), ajuste o limite e apague os gastos antigos (ou duplique a aba antes).",
               "", "Planilha gerada automaticamente pelo script gerar_planilha.py."]
        for i, t in enumerate(txt):
            c = h.cell(i + 1, 1, t)
            c.font = fonte(size=10, bold=(i == 0), color=NAVY if i == 0 else "000000")

    def gerar(self, saida):
        self.contas_mensais()
        self.gastos_do_mes()
        self.como_usar()
        self.wb.calculation = CalcProperties(fullCalcOnLoad=True)
        self.wb.save(saida)
        return saida


def main():
    pasta = Path(__file__).parent
    p = argparse.ArgumentParser(description="Gera a planilha Contas Mensais.")
    p.add_argument("--config", default=pasta / "config.json")
    p.add_argument("--dados", default=pasta / "dados" / "exemplo.json")
    p.add_argument("--saida", default=pasta / "Contas_Mensais.xlsx")
    a = p.parse_args()
    cfg = json.loads(Path(a.config).read_text(encoding="utf-8"))
    dados = json.loads(Path(a.dados).read_text(encoding="utf-8"))
    print("Planilha criada:", Gerador(cfg, dados).gerar(a.saida))


if __name__ == "__main__":
    main()
